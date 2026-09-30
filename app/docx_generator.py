"""Сборка .docx из протокола встречи (python-docx)."""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.protocol_parser import (
    META_KEYS,
    PLACEHOLDER_CELL,
    PLACEHOLDER_EMPTY,
    PLACEHOLDER_META,
    SECTION_DECISIONS,
    SECTION_DISCUSSION,
    SECTION_METADATA,
    SECTION_RISKS,
    SECTION_SUMMARY,
    SECTION_TASKS,
    TASKS_COLUMNS,
    Protocol,
    parse_protocol,
)

FONT_NAME = "Calibri"
FONT_SIZES = {"Normal": 11, "Heading 1": 16, "Heading 2": 13, "Heading 3": 11}
MARGIN_VERTICAL = Cm(2.0)
MARGIN_HORIZONTAL = Cm(1.5)
TASK_COLUMN_WIDTHS = (Cm(10.0), Cm(4.5), Cm(3.5))
BLACK = RGBColor(0, 0, 0)
DOC_TITLE_PREFIX = "Протокол встречи: "


def _set_font(style, size_pt: int) -> None:
    style.font.name = FONT_NAME
    style.font.size = Pt(size_pt)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    # Темовые шрифты заголовков перекрывают явный font.name — убираем их.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)


def _apply_styles(doc: DocumentObject) -> None:
    for style_name, size in FONT_SIZES.items():
        style = doc.styles[style_name]
        _set_font(style, size)
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.font.color.rgb = BLACK
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
    normal = doc.styles["Normal"].paragraph_format
    normal.space_before = Pt(0)
    normal.space_after = Pt(6)
    normal.line_spacing = 1.15
    for style_name in ("List Bullet", "List Number"):
        doc.styles[style_name].paragraph_format.space_after = Pt(3)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = MARGIN_VERTICAL
        section.left_margin = section.right_margin = MARGIN_HORIZONTAL


def _add_placeholder(doc: DocumentObject) -> None:
    doc.add_paragraph(PLACEHOLDER_EMPTY)


def _add_list(doc: DocumentObject, items: list[str], style: str) -> None:
    if not items:
        _add_placeholder(doc)
        return
    for item in items:
        doc.add_paragraph(item, style=style)


def _add_metadata(doc: DocumentObject, metadata: dict[str, str]) -> None:
    for key in META_KEYS:
        value = (metadata.get(key) or "").strip() or PLACEHOLDER_META
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{key}: ").bold = True
        paragraph.add_run(value)


def _add_discussion(doc: DocumentObject, discussion: list[tuple[str, str]]) -> None:
    if not discussion:
        _add_placeholder(doc)
        return
    for topic, body in discussion:
        doc.add_heading(topic, level=3)
        paragraphs = [line.strip() for line in body.split("\n") if line.strip()]
        if not paragraphs:
            _add_placeholder(doc)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)


def _mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_table_borders(table) -> None:
    # Явные границы: стиль «Table Grid» не все редакторы отображают одинаково.
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")
        borders.append(element)
    tbl_pr = table._tbl.tblPr
    # По схеме OOXML tblBorders идёт перед shd/tblLayout/tblCellMar/tblLook.
    following = tbl_pr.first_child_found_in("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    if following is None:
        tbl_pr.append(borders)
    else:
        following.addprevious(borders)


def _add_tasks(doc: DocumentObject, protocol: Protocol) -> None:
    if not protocol.tasks:
        _add_placeholder(doc)
        return
    table = doc.add_table(rows=1, cols=len(TASKS_COLUMNS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    header = table.rows[0]
    _mark_header_row(header)
    for cell, title in zip(header.cells, TASKS_COLUMNS):
        cell.text = ""
        cell.paragraphs[0].add_run(title).bold = True
    for task in protocol.tasks:
        cells = table.add_row().cells
        cells[0].text = task.task.strip()
        cells[1].text = (task.owner or "").strip() or PLACEHOLDER_CELL
        cells[2].text = (task.due or "").strip() or PLACEHOLDER_CELL
    for column, width in zip(table.columns, TASK_COLUMN_WIDTHS):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, TASK_COLUMN_WIDTHS):
            cell.width = width
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)


def build_docx(protocol: Protocol) -> DocumentObject:
    doc = Document()
    _apply_styles(doc)
    title = protocol.title.strip() or PLACEHOLDER_META
    doc.core_properties.title = f"{DOC_TITLE_PREFIX}{title}"

    doc.add_heading(f"{DOC_TITLE_PREFIX}{title}", level=1)

    doc.add_heading(SECTION_METADATA, level=2)
    _add_metadata(doc, protocol.metadata)

    doc.add_heading(SECTION_DISCUSSION, level=2)
    _add_discussion(doc, protocol.discussion)

    doc.add_heading(SECTION_DECISIONS, level=2)
    _add_list(doc, protocol.decisions, "List Number")

    doc.add_heading(SECTION_TASKS, level=2)
    _add_tasks(doc, protocol)

    doc.add_heading(SECTION_RISKS, level=2)
    _add_list(doc, protocol.risks, "List Bullet")

    if protocol.summary is not None:
        doc.add_heading(SECTION_SUMMARY, level=2)
        _add_list(doc, protocol.summary, "List Bullet")

    return doc


def protocol_to_docx_bytes(text: str) -> bytes:
    buffer = io.BytesIO()
    build_docx(parse_protocol(text)).save(buffer)
    return buffer.getvalue()


def protocol_to_docx_file(text: str, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_bytes(protocol_to_docx_bytes(text))
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Конвертирует протокол встречи (markdown) в .docx")
    parser.add_argument("input", type=Path, help="путь к .md с протоколом")
    parser.add_argument("output", type=Path, help="путь к итоговому .docx")
    args = parser.parse_args(argv)

    try:
        text = args.input.read_text(encoding="utf-8")
        out = protocol_to_docx_file(text, args.output)
    except (OSError, ValueError) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

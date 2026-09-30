import io

import pytest
from docx import Document

from app.docx_generator import build_docx, main, protocol_to_docx_bytes, protocol_to_docx_file
from app.protocol_parser import Protocol, Task, parse_protocol

H2_SECTIONS = ["Метаданные", "Обсуждение", "Решения", "Задачи", "Риски и открытые вопросы"]
TASKS_HEADER = ["Задача", "Ответственный", "Срок"]
EMPTY = "— не указано —"


def _open(data: bytes):
    return Document(io.BytesIO(data))


def _save(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _headings(doc, level: int) -> list[str]:
    return [p.text for p in doc.paragraphs if p.style.name == f"Heading {level}"]


def _texts(doc) -> list[str]:
    return [p.text for p in doc.paragraphs]


def _rows(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def _section_texts(doc, section: str) -> list[str]:
    """Тексты абзацев между заголовком H2 `section` и следующим H2."""
    result, inside = [], False
    for p in doc.paragraphs:
        if p.style.name == "Heading 2":
            inside = p.text == section
            continue
        if inside:
            result.append(p.text)
    return result


def test_full_protocol_has_title_and_all_sections(sample_protocol):
    data = protocol_to_docx_bytes(sample_protocol)
    assert data[:2] == b"PK"
    doc = _open(data)

    assert _headings(doc, 1) == ["Протокол встречи: Планирование релиза v2.0"]
    assert _headings(doc, 2) == H2_SECTIONS
    assert _headings(doc, 3) == ["Дата релиза", "Аналитика"]
    assert doc.core_properties.title == "Протокол встречи: Планирование релиза v2.0"


def test_sections_content(sample_protocol):
    doc = _open(protocol_to_docx_bytes(sample_protocol))

    assert _section_texts(doc, "Метаданные") == [
        "Дата: 14.10.2025",
        "Время: 11:00-12:00",
        "Место: Zoom",
        "Участники: Анна Смирнова (PM), Дмитрий Ковалёв (backend), Игорь Петров (QA)",
    ]
    discussion = _section_texts(doc, "Обсуждение")
    assert "Анна Смирнова зафиксировала дату 12.11 как окончательную." in discussion
    assert [p.text for p in doc.paragraphs if p.style.name == "List Number"] == [
        "Перенести релиз v2.0 на 12.11.2025.",
        "Включить пуш-уведомления в релиз v2.0.",
    ]
    assert _section_texts(doc, "Риски и открытые вопросы") == [
        "Не назначен ответственный: обновить release notes для стора.",
        "Не принято решение по системе аналитики: Amplitude или Firebase.",
    ]


def test_tasks_table_matches_input(sample_protocol):
    doc = _open(protocol_to_docx_bytes(sample_protocol))

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.columns) == 3
    rows = _rows(table)
    assert rows[0] == TASKS_HEADER
    assert rows[1:] == [
        ["Завершить разработку API пуш-уведомлений", "Дмитрий Ковалёв", "24.10.2025"],
        ["Подготовить план регрессионного тестирования", "Игорь Петров", "31.10.2025"],
        ["Обновить release notes для стора", "не указан", "не указан"],
    ]
    assert all(run.bold for cell in table.rows[0].cells for run in cell.paragraphs[0].runs)


def test_missing_fields_keep_placeholder(sample_protocol):
    rows = _rows(_open(protocol_to_docx_bytes(sample_protocol)).tables[0])
    assert rows[3][1:] == ["не указан", "не указан"]


@pytest.mark.parametrize("owner, due", [("", ""), ("   ", "  "), ("не указан", "не указан")])
def test_missing_fields_in_objects_filled_with_placeholder(owner, due):
    protocol = Protocol(title="Отчёт", tasks=[Task(task="Подготовить отчёт", owner=owner, due=due)])
    doc = _open(_save(build_docx(protocol)))
    assert _rows(doc.tables[0])[1] == ["Подготовить отчёт", "не указан", "не указан"]


def test_empty_sections_do_not_break_generation(empty_protocol):
    doc = _open(protocol_to_docx_bytes(empty_protocol))

    assert _headings(doc, 1) == ["Протокол встречи: не указано"]
    assert _headings(doc, 2) == H2_SECTIONS
    assert doc.tables == []
    for section in ("Обсуждение", "Решения", "Задачи", "Риски и открытые вопросы"):
        assert _section_texts(doc, section) == [EMPTY]
    assert "Участники: не указано" in _section_texts(doc, "Метаданные")


def test_empty_protocol_object_builds():
    doc = _open(_save(build_docx(Protocol(title=""))))

    assert _headings(doc, 1) == ["Протокол встречи: не указано"]
    assert _headings(doc, 2) == H2_SECTIONS
    assert _texts(doc).count(EMPTY) == 4
    assert _section_texts(doc, "Метаданные") == [
        "Дата: не указано",
        "Время: не указано",
        "Место: не указано",
        "Участники: не указано",
    ]


def test_parser_round_trip(sample_protocol):
    protocol = parse_protocol(sample_protocol)

    assert protocol.title == "Планирование релиза v2.0"
    assert protocol.metadata["Дата"] == "14.10.2025"
    assert protocol.decisions == [
        "Перенести релиз v2.0 на 12.11.2025.",
        "Включить пуш-уведомления в релиз v2.0.",
    ]
    assert protocol.tasks == [
        Task("Завершить разработку API пуш-уведомлений", "Дмитрий Ковалёв", "24.10.2025"),
        Task("Подготовить план регрессионного тестирования", "Игорь Петров", "31.10.2025"),
        Task("Обновить release notes для стора", "не указан", "не указан"),
    ]
    assert protocol.risks == [
        "Не назначен ответственный: обновить release notes для стора.",
        "Не принято решение по системе аналитики: Amplitude или Firebase.",
    ]
    assert protocol.discussion[1] == ("Аналитика", "Обсуждался переход на Amplitude вместо Firebase. Решение не принято.")


def test_docx_matches_parsed_protocol(sample_protocol):
    protocol = parse_protocol(sample_protocol)
    doc = _open(protocol_to_docx_bytes(sample_protocol))

    assert _rows(doc.tables[0])[1:] == [[t.task, t.owner, t.due] for t in protocol.tasks]
    assert _headings(doc, 3) == [topic for topic, _ in protocol.discussion]


def test_styles_and_margins(sample_protocol):
    doc = _open(protocol_to_docx_bytes(sample_protocol))
    for name, size in {"Normal": 11, "Heading 1": 16, "Heading 2": 13, "Heading 3": 11}.items():
        assert doc.styles[name].font.name == "Calibri"
        assert doc.styles[name].font.size.pt == size
    section = doc.sections[0]
    # Word хранит поля в твипах, поэтому сравниваем с точностью до сотых сантиметра.
    assert [round(m.cm, 2) for m in (section.top_margin, section.bottom_margin)] == [2.0, 2.0]
    assert [round(m.cm, 2) for m in (section.left_margin, section.right_margin)] == [1.5, 1.5]


def test_summary_section_rendered(sample_protocol):
    doc = _open(protocol_to_docx_bytes(sample_protocol + "\n## Краткое резюме\n\n- Релиз 12.11.2025.\n"))
    assert _headings(doc, 2) == H2_SECTIONS + ["Краткое резюме"]
    assert _section_texts(doc, "Краткое резюме") == ["Релиз 12.11.2025."]


def test_invalid_protocol_raises(sample_protocol):
    with pytest.raises(ValueError):
        protocol_to_docx_bytes(sample_protocol.replace("## Решения", "## Итоги"))


def test_file_and_cli(tmp_path, sample_protocol, capsys):
    out = protocol_to_docx_file(sample_protocol, tmp_path / "nested" / "a.docx")
    assert out.is_file()
    assert _headings(Document(str(out)), 1) == ["Протокол встречи: Планирование релиза v2.0"]

    src = tmp_path / "in.md"
    src.write_text(sample_protocol, encoding="utf-8")
    assert main([str(src), str(tmp_path / "b.docx")]) == 0
    assert (tmp_path / "b.docx").is_file()

    src.write_text("не протокол", encoding="utf-8")
    assert main([str(src), str(tmp_path / "c.docx")]) == 1
    assert "Ошибка" in capsys.readouterr().err
    assert not (tmp_path / "c.docx").exists()
